from dataclasses import dataclass
from pathlib import Path
import re
import json
import shutil
import subprocess
import tempfile
from zipfile import ZipFile

TEXT_FORMATS = {
    '.txt', '.md', '.markdown', '.rst', '.log', '.csv', '.tsv', '.json', '.jsonl',
    '.yaml', '.yml', '.toml', '.xml', '.ini', '.cfg', '.conf', '.sql',
    '.py', '.pyi', '.js', '.jsx', '.ts', '.tsx', '.mjs', '.cjs', '.java', '.kt',
    '.kts', '.c', '.h', '.cpp', '.hpp', '.cc', '.cs', '.go', '.rs', '.rb', '.php',
    '.swift', '.dart', '.r', '.sh', '.bash', '.ps1', '.bat', '.cmd', '.css',
    '.scss', '.sass', '.less', '.vue', '.svelte', '.qml', '.html', '.htm',
}
IMAGE_FORMATS = {'.png', '.jpg', '.jpeg', '.webp', '.bmp', '.tif', '.tiff', '.heic', '.heif'}
SUPPORTED = TEXT_FORMATS | IMAGE_FORMATS | {
    '.pdf', '.doc', '.docx', '.xls', '.xlsx', '.ppt', '.pptx',
    '.epub', '.odt', '.ods', '.odp', '.ipynb',
}
MAX_TEXT_BYTES = 10 * 1024 * 1024
MAX_EXTRACTED_CHARS = 2_000_000


def read_text(path):
    if path.stat().st_size > MAX_TEXT_BYTES:
        raise ValueError('Text file exceeds the 10 MB extraction limit')
    with path.open('rb') as stream:
        raw = stream.read(MAX_TEXT_BYTES + 1)
    if len(raw) > MAX_TEXT_BYTES:
        raise ValueError('Text file exceeds the 10 MB extraction limit')
    encoding = 'utf-16' if raw.startswith((b'\xff\xfe', b'\xfe\xff')) else 'utf-8-sig'
    if encoding != 'utf-16' and b'\x00' in raw:
        raise ValueError('Binary content in a text file')
    try:
        return raw.decode(encoding)
    except UnicodeDecodeError:
        return raw.decode('cp1252')


def check_office_archive(path):
    from zipfile import ZipFile
    with ZipFile(path) as archive:
        if sum(item.file_size for item in archive.infolist()) > 100 * 1024 * 1024:
            raise ValueError('Office document exceeds the 100 MB expanded-size limit')


def register_image_plugins():
    """Register optional local decoders without making them a hard import dependency."""
    try:
        from pillow_heif import register_heif_opener
    except ImportError:
        return False
    register_heif_opener()
    return True


def archive_text(path, suffix):
    """Extract readable text from EPUB/OpenDocument packages with size limits."""
    check_office_archive(path)
    from html.parser import HTMLParser

    class VisibleText(HTMLParser):
        def __init__(self):
            super().__init__(convert_charrefs=True)
            self.hidden = 0
            self.parts = []

        def handle_starttag(self, tag, attrs):
            if tag in {'script', 'style', 'noscript'}:
                self.hidden += 1

        def handle_endtag(self, tag):
            if tag in {'script', 'style', 'noscript'}:
                self.hidden = max(0, self.hidden - 1)

        def handle_data(self, data):
            if not self.hidden:
                self.parts.append(data)

    passages = []
    total = 0
    with ZipFile(path) as archive:
        names = archive.namelist()
        if suffix == '.epub':
            names = [name for name in names if name.lower().endswith(('.xhtml', '.html', '.htm'))]
        elif suffix in {'.odt', '.ods', '.odp'}:
            names = ['content.xml'] if 'content.xml' in names else []
        for name in names:
            raw = archive.read(name)
            parser = VisibleText()
            parser.feed(raw.decode('utf-8', errors='replace'))
            text = ' '.join(' '.join(parser.parts).split())
            if not text:
                continue
            total += len(text)
            if total > MAX_EXTRACTED_CHARS:
                raise ValueError('Document exceeds the extracted-text limit')
            passages.append(Passage(text))
    return passages


def notebook_passages(path):
    if path.stat().st_size > MAX_TEXT_BYTES:
        raise ValueError('Notebook exceeds the 10 MB extraction limit')
    try:
        notebook = json.loads(path.read_text(encoding='utf-8'))
    except (UnicodeDecodeError, json.JSONDecodeError) as error:
        raise ValueError(f'Invalid Jupyter notebook: {error}') from error
    passages, total = [], 0
    for index, cell in enumerate(notebook.get('cells', []), 1):
        source = ''.join(cell.get('source', []))
        outputs = []
        for output in cell.get('outputs', []):
            text = output.get('text', [])
            if isinstance(text, list):
                outputs.extend(text)
        text = f'Cell {index} ({cell.get("cell_type", "unknown")})\n{source}'
        if outputs:
            text += '\nOutput:\n' + ''.join(outputs)
        total += len(text)
        if total > MAX_EXTRACTED_CHARS:
            raise ValueError('Notebook exceeds the extracted-text limit')
        if source.strip() or outputs:
            passages.append(Passage(text))
    return passages


def legacy_office_passages(path):
    """Use an installed local LibreOffice converter for binary Office 97 files."""
    executable = shutil.which('soffice') or shutil.which('libreoffice')
    if not executable:
        raise ValueError('Legacy Office files require a local LibreOffice installation')
    with tempfile.TemporaryDirectory(prefix='providence-convert-') as directory:
        result = subprocess.run(
            [executable, '--headless', '--convert-to', 'txt:Text', '--outdir', directory, str(path)],
            capture_output=True, text=True, timeout=90, check=False,
        )
        converted = Path(directory) / f'{path.stem}.txt'
        if result.returncode or not converted.is_file():
            detail = (result.stderr or result.stdout).strip()
            raise ValueError(f'Legacy Office conversion failed{": " + detail if detail else ""}')
        return [Passage(read_text(converted))]


def spreadsheet_passages(path):
    from openpyxl import load_workbook
    check_office_archive(path)
    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    passages, count = [], 0
    try:
        for sheet in workbook:
            # Ignore unreliable worksheet dimension declarations; stream real rows.
            sheet.reset_dimensions()
            lines = [f'Sheet: {sheet.title}']
            for number, row in enumerate(sheet.iter_rows(values_only=True, max_row=100_001), 1):
                if number > 100_000 or len(row) > 1000:
                    raise ValueError('Spreadsheet exceeds the row/column extraction limit')
                text = ' | '.join('' if value is None else str(value) for value in row)
                if not text.strip(' |'):
                    continue
                line = f'Row {number}: {text}'
                count += len(line)
                if count > MAX_EXTRACTED_CHARS:
                    raise ValueError('Spreadsheet exceeds the extracted-text limit')
                lines.append(line)
                if len(lines) >= 101:
                    passages.append(Passage('\n'.join(lines)))
                    lines = [f'Sheet: {sheet.title}']
            if len(lines) > 1:
                passages.append(Passage('\n'.join(lines)))
    finally:
        workbook.close()
    return passages


@dataclass(frozen=True)
class Passage:
    text: str
    page: int | None = None
    ocr: bool = False


def extract(path: Path) -> list[Passage]:
    suffix = path.suffix.lower()
    if suffix in TEXT_FORMATS:
        content = read_text(path)
        if suffix in {'.html', '.htm'}:
            from html.parser import HTMLParser

            class VisibleText(HTMLParser):
                def __init__(self):
                    super().__init__(convert_charrefs=True)
                    self.hidden = 0
                    self.parts = []

                def handle_starttag(self, tag, attrs):
                    if tag in {'script', 'style', 'noscript'}:
                        self.hidden += 1

                def handle_endtag(self, tag):
                    if tag in {'script', 'style', 'noscript'}:
                        self.hidden = max(0, self.hidden - 1)

                def handle_data(self, data):
                    if not self.hidden:
                        self.parts.append(data)

            parser = VisibleText()
            parser.feed(content)
            content = '\n'.join(parser.parts)
        return [Passage(content)]
    if suffix == '.ipynb':
        return notebook_passages(path)
    if suffix in {'.epub', '.odt', '.ods', '.odp'}:
        return archive_text(path, suffix)
    if suffix in {'.doc', '.xls', '.ppt'}:
        return legacy_office_passages(path)
    if suffix == '.xlsx':
        return spreadsheet_passages(path)
    if suffix == '.pptx':
        from pptx import Presentation
        check_office_archive(path)
        passages, count = [], 0
        for number, slide in enumerate(Presentation(path).slides, 1):
            parts = [shape.text for shape in slide.shapes if shape.has_text_frame]
            parts += [' | '.join(cell.text for cell in row.cells)
                      for shape in slide.shapes if shape.has_table for row in shape.table.rows]
            text = f'Slide {number}\n' + '\n'.join(parts)
            count += len(text)
            if count > MAX_EXTRACTED_CHARS:
                raise ValueError('Presentation exceeds the extracted-text limit')
            passages.append(Passage(text, number))
        return passages
    if suffix == ".pdf":
        import pymupdf

        with pymupdf.open(path) as doc:
            if doc.needs_pass:
                raise ValueError("Password-protected PDF")
            return [
                Passage(page.get_text(sort=True), i + 1) for i, page in enumerate(doc)
            ]
    if suffix == ".docx":
        from docx import Document

        check_office_archive(path)
        document = Document(path)
        texts = [paragraph.text for paragraph in document.paragraphs]
        texts.extend(
            " | ".join(cell.text for cell in row.cells)
            for table in document.tables
            for row in table.rows
        )
        return [Passage("\n".join(texts))]
    if suffix in IMAGE_FORMATS:
        register_image_plugins()
        return []
    raise ValueError(f"Unsupported file type: {suffix}")


def chunk(passages: list[Passage], words: int = 90, overlap: int = 20) -> list[Passage]:
    if not 0 <= overlap < words:
        raise ValueError("Overlap must be smaller than chunk size")
    result = []
    for passage in passages:
        tokens = re.findall(r"\S+", passage.text)
        for start in range(0, len(tokens), words - overlap):
            result.append(
                Passage(
                    " ".join(tokens[start : start + words]), passage.page, passage.ocr
                )
            )
            if start + words >= len(tokens):
                break
    return result
