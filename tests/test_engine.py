from pathlib import Path
import math

import pytest

from providence.documents import Passage, chunk, extract
from providence.engine import Engine, normalized
from providence.store import Store


class TestEmbedder:
    __test__ = False
    identity = "test-model"

    def encode(self, texts):
        return [
            [1 + text.lower().count("database"), 1 + text.lower().count("beach")]
            for text in texts
        ]


@pytest.fixture
def engine(tmp_path):
    return Engine(
        Store(tmp_path / "data" / "index.sqlite3", TestEmbedder.identity),
        TestEmbedder(),
    )


def test_incremental_index_refresh_and_delete(engine, tmp_path):
    root = tmp_path / "documents"
    root.mkdir()
    path = root / "final_v3.txt"
    path.write_text("database " * 200, encoding="utf-8")
    assert engine.index(root).indexed == 1
    assert engine.index(root).unchanged == 1
    assert len(engine.search("database")) == 1
    path.write_text("beach vacation", encoding="utf-8")
    assert engine.index(root).indexed == 1
    assert engine.search("beach")[0].text == "beach vacation"
    path.unlink()
    assert engine.index(root).removed == 1
    assert engine.search("beach") == []


def test_ranking_and_persistence(engine, tmp_path):
    root = tmp_path / "documents"
    root.mkdir()
    (root / "proposal.txt").write_text("database database schema")
    (root / "trip.txt").write_text("beach beach vacation")
    engine.index(root)
    reopened = Engine(Store(engine.store.path, TestEmbedder.identity), TestEmbedder())
    assert Path(reopened.search("database")[0].path).name == "proposal.txt"
    assert len(reopened.search("beach", limit=1)) == 1
    assert reopened.store.roots() == [root.resolve()]


def test_bad_file_does_not_abort_other_files(engine, tmp_path):
    (tmp_path / "bad.docx").write_bytes(b"invalid zip")
    (tmp_path / "good.txt").write_text("database proposal")
    report = engine.index(tmp_path)
    assert report.indexed == 1
    assert len(report.errors) == 1
    assert len(engine.search("database")) == 1


def test_empty_refresh_removes_stale_snippets(engine, tmp_path):
    path = tmp_path / "document.txt"
    path.write_text("database")
    engine.index(tmp_path)
    path.write_text("")
    assert engine.index(tmp_path).errors
    assert engine.search("database") == []


def test_unavailable_root_preserves_index(engine, tmp_path):
    root = tmp_path / "documents"
    root.mkdir()
    (root / "a.txt").write_text("database")
    engine.index(root)
    with pytest.raises(ValueError, match="unavailable"):
        engine.index(tmp_path / "missing")
    assert len(engine.search("database")) == 1


def test_embedding_failure_keeps_previous_document(engine, tmp_path):
    path = tmp_path / "a.txt"
    path.write_text("database")
    engine.index(tmp_path)
    path.write_text("beach")

    class Broken(TestEmbedder):
        def encode(self, texts):
            raise RuntimeError("missing model")

    with pytest.raises(RuntimeError, match="missing model"):
        Engine(engine.store, Broken()).index(tmp_path)
    assert engine.search("database")[0].text == "database"


def test_atomic_vector_count_validation(engine, tmp_path):
    path = tmp_path / "a.txt"
    path.write_text("database")
    engine.index(tmp_path)
    with pytest.raises(ValueError):
        engine.store.replace(path, "new", [Passage("beach")], [])
    assert engine.search("database")[0].text == "database"


def test_index_model_identity_is_enforced(engine):
    with pytest.raises(RuntimeError, match="different embedding model"):
        Store(engine.store.path, "other-model")


def test_chunk_overlap_page_and_tail():
    passages = chunk(
        [Passage(" ".join(str(i) for i in range(11)), 3)], words=5, overlap=2
    )
    assert [p.text for p in passages] == ["0 1 2 3 4", "3 4 5 6 7", "6 7 8 9 10"]
    assert all(p.page == 3 for p in passages)
    assert chunk([Passage(" \n ")]) == []
    with pytest.raises(ValueError):
        chunk([], words=2, overlap=2)


def test_txt_encodings(tmp_path):
    path = tmp_path / "text.txt"
    for encoding in ("utf-8-sig", "utf-16", "cp1252"):
        path.write_bytes("Café proposal".encode(encoding))
        assert extract(path)[0].text == "Café proposal"


def test_pdf_pages_and_docx_tables(tmp_path):
    import pymupdf
    from docx import Document

    pdf = tmp_path / "proposal.pdf"
    with pymupdf.open() as document:
        document.new_page().insert_text((72, 72), "Database proposal")
        document.new_page().insert_text((72, 72), "Implementation schedule")
        document.save(pdf)
    passages = extract(pdf)
    assert passages[0].page == 1 and "Database" in passages[0].text
    assert passages[1].page == 2 and "schedule" in passages[1].text
    docx = tmp_path / "proposal.docx"
    document = Document()
    document.add_paragraph("Library schema")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Borrowers"
    document.save(docx)
    assert "Library schema" in extract(docx)[0].text
    assert "Borrowers" in extract(docx)[0].text


def test_normalized_rejects_bad_vectors():
    for vector in ([0, 0], [math.nan, 1], [math.inf, 1]):
        with pytest.raises(ValueError):
            normalized(vector)


def test_scan_error_does_not_prune(engine, tmp_path, monkeypatch):
    import providence.engine as module

    (tmp_path / "a.txt").write_text("database")
    engine.index(tmp_path)

    def inaccessible(root, onerror, followlinks):
        onerror(PermissionError("Access denied"))
        return iter([])

    monkeypatch.setattr(module.os, "walk", inaccessible)
    assert engine.index(tmp_path).errors
    assert engine.store.count() == 1


def test_reindex_forces_embeddings(engine, tmp_path):
    (tmp_path / "a.txt").write_text("database")
    engine.index(tmp_path)
    assert engine.index(tmp_path, force=True).indexed == 1


def test_corrupt_refresh_removes_stale_text(engine, tmp_path):
    from docx import Document

    path = tmp_path / "a.docx"
    document = Document()
    document.add_paragraph("database proposal")
    document.save(path)
    engine.index(tmp_path)
    path.write_bytes(b"broken")
    assert engine.index(tmp_path).errors
    assert engine.search("database") == []


def test_pruning_one_root_preserves_another(engine, tmp_path):
    roots = [tmp_path / "first", tmp_path / "second"]
    for root in roots:
        root.mkdir()
        (root / "a.txt").write_text("database")
        engine.index(root)
    (roots[0] / "a.txt").unlink()
    assert engine.index(roots[0]).removed == 1
    assert engine.store.count() == 1


def test_symlinks_are_not_followed(engine, tmp_path):
    root = tmp_path / "documents"
    root.mkdir()
    outside = tmp_path / "outside.txt"
    outside.write_text("database")
    try:
        (root / "link.txt").symlink_to(outside)
    except OSError as error:
        if getattr(error, "winerror", None) == 1314:
            pytest.skip(
                "This Windows account does not have symlink creation privileges"
            )
        raise
    assert engine.index(root).indexed == 0
    assert engine.store.count() == 0


def test_remove_parent_preserves_selected_child(engine, tmp_path):
    parent = tmp_path / "parent"
    child = parent / "child"
    child.mkdir(parents=True)
    (parent / "a.txt").write_text("database")
    (child / "b.txt").write_text("database")
    engine.index(parent)
    engine.index(child)
    engine.store.remove_root(parent)
    assert engine.store.roots() == [child.resolve()]
    assert engine.store.count() == 1
    assert engine.search("database")[0].path.endswith("b.txt")


def test_many_small_files_share_batches_and_progress_is_monotonic(tmp_path):
    from providence.engine import IndexProgress

    class Counting(TestEmbedder):
        def __init__(self):
            self.batches = []

        def encode(self, texts):
            self.batches.append(len(texts))
            return super().encode(texts)

    embedder = Counting()
    engine = Engine(Store(tmp_path / "index.sqlite3", embedder.identity), embedder)
    for index in range(65):
        (tmp_path / f"file-{index:03}.txt").write_text(
            f"database proposal number {index}"
        )
    (tmp_path / "corrupt.docx").write_bytes(b"not a docx")
    updates = []
    report = engine.index(tmp_path, updates.append)
    assert report.indexed == 65
    assert len(report.errors) == 1
    assert embedder.batches == [32, 32, 1]
    assert engine.store.count() == 65
    assert all(isinstance(update, IndexProgress) for update in updates)
    assert [update.percent for update in updates] == sorted(
        update.percent for update in updates
    )
    assert updates[-1].percent == 100
    assert updates[-1].completed == updates[-1].total == 66
    assert engine.search("database proposal number 64", limit=100)


def test_unchanged_metadata_skips_reads_extraction_and_model(
    engine, tmp_path, monkeypatch
):
    import providence.engine as module

    path = tmp_path / "note.md"
    path.write_text("# Database\nMarkdown proposal")
    engine.index(tmp_path)

    def forbidden(*args, **kwargs):
        raise AssertionError("Unchanged documents should not be read or embedded")

    monkeypatch.setattr(module, "content_digest", forbidden)
    monkeypatch.setattr(module, "extract", forbidden)
    monkeypatch.setattr(engine.embedder, "encode", forbidden)
    assert engine.index(tmp_path).unchanged == 1


def test_touched_file_reuses_hash_and_changed_file_reembeds(
    engine, tmp_path, monkeypatch
):
    import os

    path = tmp_path / "note.txt"
    path.write_text("database")
    engine.index(tmp_path)
    before = path.stat()
    os.utime(path, ns=(before.st_atime_ns, before.st_mtime_ns + 1_000_000_000))
    assert engine.index(tmp_path).unchanged == 1
    path.write_text("beach vacation")
    assert engine.index(tmp_path).indexed == 1
    assert engine.search("beach")[0].text == "beach vacation"


def test_allowed_formats_only_and_nested_roots_are_deduplicated(engine, tmp_path):
    parent = tmp_path / "documents"
    child = parent / "project"
    child.mkdir(parents=True)
    (parent / "README.MD").write_text("# Database proposal")
    (child / "note.txt").write_text("database implementation")
    for suffix in ("doc", "jpg", "exe", "csv", "html"):
        (parent / f"ignored.{suffix}").write_text("database")
    reports = engine.index_many([child, parent])
    assert sum(report.indexed for report in reports) == 2
    assert engine.store.count() == 2
    assert engine.store.roots() == sorted([parent.resolve(), child.resolve()])
    assert all(
        Path(path).suffix.lower() in {".txt", ".md"}
        for path in engine.store.documents()
    )


def test_old_index_schema_migrates_without_losing_documents(tmp_path):
    import sqlite3

    database = tmp_path / "old.sqlite3"
    with sqlite3.connect(database) as db:
        db.execute(
            "CREATE TABLE documents (path TEXT PRIMARY KEY, digest TEXT NOT NULL)"
        )
        db.execute("INSERT INTO documents VALUES ('saved.txt','original')")
    store = Store(database, TestEmbedder.identity)
    assert store.digest("saved.txt") == "original"
    assert store.snapshots()["saved.txt"] == ("original", None, None, None)


def test_stop_keeps_completed_files_and_resume_skips_them(tmp_path):
    from threading import Event

    stop = Event()

    class Stopping(TestEmbedder):
        def encode(self, texts):
            vectors = super().encode(texts)
            stop.set()
            return vectors

    embedder = Stopping()
    engine = Engine(Store(tmp_path / "index.sqlite3", embedder.identity), embedder)
    for index in range(65):
        (tmp_path / f"file-{index:03}.txt").write_text(f"database proposal {index}")
    updates = []
    report = engine.index(tmp_path, updates.append, cancelled=stop.is_set)
    assert report.cancelled
    assert report.indexed == engine.store.count() == 32
    assert updates[-1].startswith("Stopped")
    assert updates[-1].percent < 100
    assert len(Engine(engine.store, TestEmbedder()).search("database", limit=100)) == 32
    resumed = Engine(engine.store, TestEmbedder()).index(tmp_path)
    assert not resumed.cancelled
    assert resumed.unchanged == 32
    assert resumed.indexed == 33
    assert engine.store.count() == 65


def test_stop_during_scan_does_not_prune_existing_files(engine, tmp_path):
    from threading import Event

    path = tmp_path / "a.txt"
    path.write_text("database proposal")
    engine.index(tmp_path)
    path.unlink()
    stop = Event()
    stop.set()
    report = engine.index(tmp_path, cancelled=stop.is_set)
    assert report.cancelled
    assert report.removed == 0
    assert engine.store.count() == 1


def test_stop_does_not_save_partially_embedded_document(tmp_path):
    from threading import Event

    stop = Event()

    class Stopping(TestEmbedder):
        def encode(self, texts):
            stop.set()
            return super().encode(texts)

    embedder = Stopping()
    engine = Engine(Store(tmp_path / "index.sqlite3", embedder.identity), embedder)
    (tmp_path / "long.txt").write_text("database " * 5000)
    report = engine.index(tmp_path, cancelled=stop.is_set)
    assert report.cancelled
    assert report.indexed == 0
    assert engine.store.count() == 0


def test_batched_vectors_stay_with_their_document(engine, tmp_path):
    texts = ["database " * 300, "beach beach vacation", "database beach planning"]
    for index, text in enumerate(texts):
        (tmp_path / f"note-{index}.txt").write_text(text)
    assert engine.index(tmp_path).indexed == 3
    for path, text, page, vector in engine.store.vectors():
        expected = normalized(TestEmbedder().encode([text])[0])
        assert list(vector) == pytest.approx(expected)
        assert text in texts[int(Path(path).stem.split("-")[1])]
