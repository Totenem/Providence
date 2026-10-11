import json
from pathlib import Path

import pytest
from providence.documents import extract, SUPPORTED
from providence.library_tools import ImageLibrary, find_duplicates, signature
from providence.store import Store


def catalog(tmp_path, files):
    root = tmp_path / 'documents'
    root.mkdir(exist_ok=True)
    store = Store(tmp_path / 'data' / 'index.sqlite3', 'test')
    store.add_root(root)
    for name, data in files.items():
        path = root / name
        path.write_bytes(data)
        store.discover([(path, signature(path), path.stat().st_mtime, path.suffix in SUPPORTED)])
    return store, root


def test_duplicates_read_only_different_names_same_size_and_scope(tmp_path):
    store, root = catalog(tmp_path, {'a.txt': b'alpha', 'b.bin': b'alpha', 'c.txt': b'bravo'})
    before = {p: p.read_bytes() for p in root.iterdir()}
    report = find_duplicates(store)
    assert len(report.groups) == 1
    assert set(report.groups[0]['paths']) == {str(root / 'a.txt'), str(root / 'b.bin')}
    assert {p: p.read_bytes() for p in root.iterdir()} == before
    assert find_duplicates(store, root / 'elsewhere').groups == []
    assert find_duplicates(store, cancelled=lambda: True).cancelled
    store.set_setting('exclusions', json.dumps(['b.bin']))
    assert find_duplicates(store).groups == []


def test_duplicates_skip_hard_links_and_disappeared_files(tmp_path):
    store, root = catalog(tmp_path, {'a.txt': b'alpha', 'b.txt': b'alpha'})
    (root / 'b.txt').unlink()
    (root / 'b.txt').hardlink_to(root / 'a.txt')
    assert find_duplicates(store).groups == []
    (root / 'b.txt').unlink()
    assert find_duplicates(store).groups == []


class VisualModel:
    identity = 'fake-vision-v1'

    def __init__(self):
        self.calls = 0

    def image(self, path):
        self.calls += 1
        return [1., 0.] if path.name == 'beach.png' else [0., 1.]

    def text(self, query):
        return [1., 0.]


def test_images_incremental_stale_scope_and_model_isolation(tmp_path):
    store, root = catalog(tmp_path, {'beach.png': b'fake beach', 'city.png': b'fake city'})
    model = VisualModel()
    library = ImageLibrary(store, model)
    assert library.update()['indexed'] == 2
    assert library.update()['unchanged'] == 2
    assert model.calls == 2
    assert Path(library.search('beach')[0].path).name == 'beach.png'
    assert library.search('beach', root / 'elsewhere') == []
    (root / 'beach.png').write_bytes(b'changed beach')
    assert len(library.search('beach')) == 1
    assert library.update()['indexed'] == 1
    model.identity = 'new-model'
    assert library.search('beach') == []
    assert library.update()['indexed'] == 2
    (root / 'city.png').unlink()
    assert len(library.search('beach')) == 1
    assert library.search('beach', cancelled=lambda: True) == []
    assert library.update(cancelled=lambda: True)['cancelled']
    store.remove_root(root)
    assert library.search('beach') == []


@pytest.mark.parametrize('extension', ['py', 'tsx', 'rs', 'csv', 'tsv', 'json', 'yaml', 'toml', 'xml'])
def test_text_formats_and_encoding(tmp_path, extension):
    path = tmp_path / f'document.{extension}'
    path.write_text('café database schema', encoding='utf-16')
    assert extract(path)[0].text == 'café database schema'
    path.write_bytes(b'bad\x00binary')
    with pytest.raises(ValueError, match='Binary'):
        extract(path)


def test_html_reads_text_without_scripts_or_remote_resources(tmp_path):
    path = tmp_path / 'page.html'
    path.write_text('<style>hidden</style><h1>Database &amp; schema</h1><script>bad()</script><img src="https://invalid.example/image">')
    assert extract(path)[0].text == 'Database & schema'


def test_xlsx_and_pptx_passages(tmp_path):
    from openpyxl import Workbook
    from pptx import Presentation
    workbook = Workbook()
    workbook.active.title = 'Accounts'
    workbook.active.append(['Invoice', 'Amount'])
    workbook.active.append(['ZX92847', 125])
    path = tmp_path / 'accounts.xlsx'
    workbook.save(path)
    passage = extract(path)[0]
    assert 'Sheet: Accounts' in passage.text and 'Row 2: ZX92847 | 125' in passage.text
    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[1]).shapes.title.text = 'Database proposal'
    path = tmp_path / 'proposal.pptx'
    deck.save(path)
    assert extract(path)[0].page == 1
    assert 'Database proposal' in extract(path)[0].text


def test_new_support_requeues_previously_unsupported_file(tmp_path):
    store, root = catalog(tmp_path, {'data.csv': b'id,name\n1,example'})
    path = root / 'data.csv'
    store.state(path, 'unsupported')
    store.discover([(path, signature(path), path.stat().st_mtime, True)])
    assert path in store.pending()
