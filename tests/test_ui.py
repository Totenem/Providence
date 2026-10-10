import time
from pathlib import Path
from threading import Event
import pytest
from PySide6.QtWidgets import QApplication
from providence.engine import Engine
from providence.store import Store
from providence.ui import Window
from test_engine import TestEmbedder


@pytest.fixture
def app():
    return QApplication.instance() or QApplication([])


def finish(app, window, timeout=15):
    deadline = time.monotonic() + timeout
    while (window.task or window.search_task) and time.monotonic() < deadline:
        app.processEvents()
        time.sleep(0.01)
    assert window.task is None and window.search_task is None


def make_engine(tmp_path):
    return Engine(
        Store(tmp_path / "data" / "index.sqlite3", TestEmbedder.identity),
        TestEmbedder(),
    )


def library(tmp_path):
    root = tmp_path / "documents"
    root.mkdir()
    (root / "final_v3.txt").write_text("database proposal\nFull document ending")
    return root


def test_onboarding_and_search_remain_accessible(app, tmp_path):
    engine = make_engine(tmp_path)
    root = library(tmp_path)
    (tmp_path / "private.txt").write_text("private notes")
    window = Window(engine, monitor=False)
    window.show()
    assert window.pages.currentIndex() == 0 and not window.home_nav.isEnabled()
    window.add_pending_folder(root)
    window.add_pending_folder(root)
    assert window.onboarding_folders.count() == 1
    window.start_onboarding()
    assert window.pages.currentIndex() == 1 and window.query.isEnabled()
    finish(app, window)
    assert engine.store.count() == 1
    assert engine.store.roots() == [root.resolve()]
    window.query.setText("database")
    window.search()
    finish(app, window)
    assert window.file_model.rowCount() == 1
    assert "database proposal" in window.preview.toPlainText()
    assert window.open_button.isEnabled()
    assert window.preview.paper.extraSelections()
    window.close()
    reopened = Window(engine, monitor=False)
    assert reopened.pages.currentIndex() == 1
    reopened.close()


def test_invalid_location_is_recoverable(app, tmp_path):
    window = Window(make_engine(tmp_path), monitor=False)
    window.add_pending_folder(tmp_path / "missing")
    window.start_onboarding()
    assert not window.onboarded and window.pages.currentIndex() == 0
    assert "unavailable" in window.onboarding_feedback.text()
    window.close()


def test_source_change_requires_review(app, tmp_path, monkeypatch):
    engine = make_engine(tmp_path)
    engine.store.set_setting("onboarding_complete", "1")
    engine.store.set_setting("library_source", "old")
    monkeypatch.setenv("PROVIDENCE_LIBRARY_SOURCE", "new")
    window = Window(engine, monitor=False)
    assert not window.onboarded
    window.close()


def test_controls_work_during_indexing_and_close_cancels(app, tmp_path):
    root = library(tmp_path)
    entered, release = Event(), Event()

    class Slow(TestEmbedder):
        def encode(self, texts):
            entered.set()
            assert release.wait(8)
            return super().encode(texts)

    engine = Engine(
        Store(tmp_path / "data" / "index.sqlite3", TestEmbedder.identity), Slow()
    )
    window = Window(engine, monitor=False)
    window.add_pending_folder(root)
    window.start_onboarding()
    deadline = time.monotonic() + 5
    while not entered.is_set() and time.monotonic() < deadline:
        app.processEvents()
        time.sleep(0.01)
    assert entered.is_set()
    assert window.query.isEnabled() and window.folder_nav.isEnabled()
    window.navigate(2)
    assert window.pages.currentIndex() == 2
    window.browse_library()
    assert window.file_model.rowCount() == 1
    assert engine.store.file(root / "final_v3.txt")["state"] == "processing"
    window.stop_indexing()
    release.set()
    finish(app, window)
    assert window.index_cancel.is_set()
    window.close()


def test_filters_pagination_navigation_and_context_actions(app, tmp_path, monkeypatch):
    root = library(tmp_path)
    child = root / "project"
    child.mkdir()
    for i in range(210):
        (child / f"file-{i:03}.bin").write_bytes(b"hello")
    engine = make_engine(tmp_path)
    engine.index(root)
    engine.store.set_setting("onboarding_complete", "1")
    window = Window(engine, monitor=False)
    assert window.total == 211 and window.file_model.rowCount() == 200
    window.page_move(1)
    assert window.file_model.rowCount() == 11
    window.state_filter.setCurrentIndex(window.state_filter.findData("ready"))
    assert window.total == 1
    window.clear_filters()
    window.browse_library(child)
    assert window.total == 210
    window.go_up()
    assert window.active_root == root
    window.history_move(-1)
    assert window.active_root == child
    window.browse_library(root)
    window.type_filter.setCurrentIndex(window.type_filter.findData(".txt"))
    assert window.total == 1
    opened = []
    monkeypatch.setattr(
        "providence.ui.QDesktopServices.openUrl",
        lambda url: opened.append(url.toLocalFile()) or True,
    )
    window.open_document()
    assert Path(opened[0]) == root / "final_v3.txt"
    window.copy_path()
    assert QApplication.clipboard().text().endswith("final_v3.txt")
    (root / "final_v3.txt").unlink()
    window.open_document()
    assert "moved or been deleted" in window.status.text()
    window.close()


def test_theme_motion_and_layout_persist(app, tmp_path):
    engine = make_engine(tmp_path)
    root = library(tmp_path)
    engine.index(root)
    engine.store.set_setting("onboarding_complete", "1")
    window = Window(engine, monitor=False)
    window.theme_combo.setCurrentText("Dark")
    window.reduced_motion.setChecked(True)
    window.results.setColumnWidth(0, 321)
    window.navigate(3)
    assert "#111821" in window.styleSheet()
    window.close()
    reopened = Window(engine, monitor=False)
    assert reopened.theme_combo.currentText() == "Dark"
    assert reopened.reduced_motion.isChecked()
    assert reopened.results.columnWidth(0) == 321
    reopened.close()


def test_scope_removal_preserves_originals(app, tmp_path):
    engine = make_engine(tmp_path)
    root = library(tmp_path)
    engine.index(root)
    engine.store.set_setting("onboarding_complete", "1")
    window = Window(engine, monitor=False)
    window.folders.setCurrentRow(0)
    window.remove_folder()
    assert engine.store.roots() == [] and engine.store.stats() == {}
    assert (root / "final_v3.txt").is_file()
    window.close()


def test_stale_search_does_not_replace_new_scope(app, tmp_path):
    root = library(tmp_path)
    engine = make_engine(tmp_path)
    engine.index(root)
    engine.store.set_setting("onboarding_complete", "1")
    window = Window(engine, monitor=False)
    window.query.setText("database")
    window.search()
    window.browse_library()
    finish(app, window)
    assert window.search_results is None and window.file_model.rowCount() == 1
    window.close()


def test_pdf_word_markdown_and_image_previews(app, tmp_path):
    import pymupdf
    from docx import Document
    from PIL import Image
    from providence.preview import DocumentPreview

    pdf = tmp_path / "two-pages.pdf"
    with pymupdf.open() as document:
        for text in ("First page", "Second page"):
            document.new_page().insert_text((72, 72), text)
        document.save(pdf)
    preview = DocumentPreview()
    preview.resize(640, 700)
    preview.load(pdf, page=2)
    assert preview.stack.currentIndex() == 2 and preview.page == 1
    assert not preview.image.pixmap().isNull()
    preview.set_match("Second")
    assert preview.match == "Second"
    preview.turn(-1)
    assert preview.position.text() == "Page 1 of 2"
    preview.find("Second")
    assert preview.page == 1
    preview.change_zoom(0.2)
    assert preview.zoom == 1.2
    word = tmp_path / "formatted.docx"
    document = Document()
    document.add_heading("Proposal", 1)
    document.add_paragraph().add_run("Bold statement").bold = True
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Table content"
    document.save(word)
    preview.load(word)
    assert "Table content" in preview.toPlainText()
    assert "font-weight:700" in preview.paper.toHtml()
    md = tmp_path / "note.md"
    md.write_text("# Heading\n\n**Bold text**")
    preview.load(md)
    assert "**" not in preview.toPlainText()
    image = tmp_path / "picture.png"
    Image.new("RGB", (100, 100), "white").save(image)
    preview.load(image)
    assert preview.stack.currentIndex() == 2
    preview.load(tmp_path / "missing.pdf")
    assert "moved or been deleted" in preview.empty.text()
    preview.close()


def test_live_watcher_updates_ui(app, tmp_path):
    engine = make_engine(tmp_path)
    root = library(tmp_path)
    engine.index(root)
    engine.store.set_setting("onboarding_complete", "1")
    window = Window(engine)
    window.show()

    def until(predicate):
        deadline = time.monotonic() + 12
        while not predicate() and time.monotonic() < deadline:
            app.processEvents()
            time.sleep(0.03)
        assert predicate()

    until(lambda: window.service.observer is not None)
    added = root / "added.txt"
    added.write_text("beach vacation")
    until(
        lambda: (
            engine.store.file(added) is not None
            and engine.store.file(added)["state"] == "ready"
        )
    )
    until(lambda: window.file_model.rowCount() == 2)
    added.unlink()
    until(lambda: engine.store.file(added) is None)
    window.close()
    until(lambda: not window.isVisible())
